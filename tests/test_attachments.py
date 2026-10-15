"""1.4.9: attach pictures and files in chats (PC window + phone app)."""
import io
import json
import os
import sys
import tempfile
import threading
import unittest
import urllib.parse
import urllib.request
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import attachments  # noqa: E402
try:
    import agent_core as agent  # noqa: E402
except ImportError:
    import agent_core as agent  # noqa: E402
import rules  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]


def png_bytes(size=(64, 48)):
    from PIL import Image
    out = io.BytesIO(); Image.new("RGB", size, (200, 30, 30)).save(out, "PNG"); return out.getvalue()


class AttachmentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        attachments.upload_dir_override = Path(self.tmp.name) / "uploads"

    def tearDown(self):
        attachments.upload_dir_override = None
        self.tmp.cleanup()

    def test_save_image_and_reject_bad_files(self):
        r = attachments.save_upload("photo.png", png_bytes())
        self.assertTrue(r["ok"], r); self.assertEqual(r["kind"], "image")
        self.assertTrue(Path(r["path"]).is_file())
        self.assertIn(attachments.upload_dir().resolve(), Path(r["path"]).resolve().parents)
        self.assertFalse(attachments.save_upload("evil.exe", b"MZ")["ok"])
        self.assertFalse(attachments.save_upload("run.bat", b"echo")["ok"])
        self.assertFalse(attachments.save_upload("fake.png", b"not a picture")["ok"])
        self.assertFalse(attachments.save_upload("big.txt", b"x" * (attachments.MAX_BYTES + 1))["ok"])
        self.assertFalse(attachments.save_upload("empty.txt", b"")["ok"])

    def test_names_cannot_escape_the_folder(self):
        r = attachments.save_upload("..\\..\\Windows\\win.ini.txt", b"hello")
        self.assertTrue(r["ok"]); self.assertEqual(Path(r["path"]).parent.resolve(), attachments.upload_dir().resolve())
        self.assertIsNone(attachments.get("../../etc"))
        self.assertIsNone(attachments.get("0" * 16))

    def test_build_turn_images_pdf_as_parts_docs_as_text(self):
        import docx
        import openpyxl
        d = docx.Document(); d.add_paragraph("Quarterly sales went up 12 percent.")
        b = io.BytesIO(); d.save(b)
        wb = openpyxl.Workbook(); wb.active.append(["Region", "Sales"]); wb.active.append(["North", 420])
        x = io.BytesIO(); wb.save(x)
        ids = [attachments.save_upload("chart.png", png_bytes((3000, 2000)))["id"],
               attachments.save_upload("report.pdf", b"%PDF-1.4 fake")["id"],
               attachments.save_upload("notes.docx", b.getvalue())["id"],
               attachments.save_upload("sales.xlsx", x.getvalue())["id"],
               attachments.save_upload("todo.txt", "buy milk".encode())["id"]]
        text, parts = attachments.build_turn("What is this?", ids)
        self.assertEqual([m for _, m in parts], ["image/jpeg", "application/pdf"])
        from PIL import Image
        self.assertLessEqual(max(Image.open(io.BytesIO(parts[0][0])).size), attachments.MAX_SIDE)
        for want in ("What is this?", "Quarterly sales went up 12 percent.", "North,420", "buy milk", "chart.png", "report.pdf"):
            self.assertIn(want, text)

    def test_marker_roundtrip_and_info(self):
        fid = attachments.save_upload("a.txt", b"hi")["id"]
        stored = attachments.message_with_marker("look", [fid, "deadbeefdeadbeef"])
        text, ids = attachments.split_marker(stored)
        self.assertEqual((text, ids), ("look", [fid]))
        self.assertEqual(attachments.info(ids)[0]["name"], "a.txt")
        self.assertEqual(attachments.message_with_marker("plain", []), "plain")

    def test_agent_sends_attached_parts_once(self):
        from google.genai import types
        agent.turn_parts = [(png_bytes(), "image/png")]
        try:
            parts = agent._user_parts(types, "what is in the picture")
            self.assertEqual(len(parts), 2)
            self.assertEqual(parts[1].inline_data.mime_type, "image/png")
            self.assertEqual(len(agent._user_parts(types, agent.RESUME_TEXT)), 1)
        finally:
            agent.turn_parts = []
        self.assertEqual(len(agent._user_parts(types, "x")), 1)


class AttachmentHttpTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        root = Path(self.tmp.name)
        (root / "rules.json").write_text("{}", encoding="utf-8")
        self.env = mock.patch.dict(os.environ, {"RULES_FILE": str(root / "rules.json"), "STORAGE_DIR": str(root / "s"),
                                                "BOTS_DIR": str(root / "b"), "GROUPS_DIR": str(root / "g")})
        self.env.start(); rules.reset_cache()
        attachments.upload_dir_override = root / "uploads"

    def tearDown(self):
        attachments.upload_dir_override = None
        self.env.stop(); rules.reset_cache()
        import chat_store
        chat_store.close(); self.tmp.cleanup()

    def test_upload_chat_history_and_download(self):
        import chat_app
        server, state = chat_app.serve("m", 0)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{server.server_address[1]}"
        seen = {}
        def fake_run(text, history, model, **kw):
            seen["text"] = text; seen["parts"] = list(agent.turn_parts); return "It is a red picture."
        try:
            def up(name, data):
                req = urllib.request.Request(base + "/api/upload", data, {"Content-Type": "application/octet-stream",
                                             "X-File-Name": urllib.parse.quote(name), "Origin": base})
                return json.loads(urllib.request.urlopen(req, timeout=5).read())
            img = up("red dot.png", png_bytes()); doc = up("plan.txt", b"step one: relax")
            self.assertTrue(img["ok"] and doc["ok"], (img, doc))
            bad = up("x.exe", b"MZ"); self.assertFalse(bad["ok"])
            events = []
            with mock.patch.object(agent, "run_agent", fake_run):
                chat_app.stream_turn(state, "att1", "what is this?", events.append, speak=False, files=[img["id"], doc["id"]])
            self.assertEqual([m for _, m in seen["parts"]], ["image/png"])
            self.assertIn("step one: relax", seen["text"])
            self.assertEqual(agent.turn_parts, [])
            rows = json.loads(urllib.request.urlopen(base + "/api/history?session=att1", timeout=5).read())
            user = [r for r in rows if r["role"] == "user"][0]
            self.assertEqual(user["text"], "what is this?")
            self.assertEqual([f["name"] for f in user["files"]], ["red dot.png", "plan.txt"])
            sess = json.loads(urllib.request.urlopen(base + "/api/sessions", timeout=5).read())
            self.assertEqual([x["title"] for x in sess if x["session"] == "att1"], ["what is this?"])
            got = urllib.request.urlopen(base + img["url"], timeout=5)
            self.assertEqual(got.headers["Content-Type"], "image/png"); self.assertEqual(got.read(), png_bytes())
            got = urllib.request.urlopen(base + doc["url"], timeout=5)
            self.assertIn("attachment", got.headers["Content-Disposition"])
            with self.assertRaises(urllib.error.HTTPError):
                urllib.request.urlopen(base + "/api/upload?id=0123456789abcdef", timeout=5)
        finally:
            chat_app.shutdown(server)


class PhoneUploadTests(unittest.TestCase):
    def test_phone_upload_needs_sign_in_then_works(self):
        import remote_access as ra
        d = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.addCleanup(d.cleanup)
        attachments.upload_dir_override = Path(d.name) / "uploads"
        self.addCleanup(setattr, attachments, "upload_dir_override", None)
        with mock.patch.object(ra, "_storage", return_value=Path(d.name)):
            ra.reset_state()
            srv = ra.ThreadingHTTPServer(("127.0.0.1", 0), ra.make_remote_handler(mock.Mock()))
            threading.Thread(target=srv.serve_forever, daemon=True).start(); self.addCleanup(srv.shutdown)
            base = f"http://127.0.0.1:{srv.server_address[1]}"
            def up(cookie):
                h = {"Content-Type": "application/octet-stream", "X-File-Name": "phone.png", "Origin": base}
                if cookie:
                    h["Cookie"] = f"{ra.COOKIE}={cookie}"
                try:
                    with urllib.request.urlopen(urllib.request.Request(base + "/api/upload", png_bytes(), h), timeout=5) as r:
                        return r.status, json.loads(r.read())
                except urllib.error.HTTPError as e:
                    return e.code, {}
            self.assertEqual(up("")[0], 401)
            self.assertFalse(list(attachments.upload_dir().glob("*.png")))
            st, body = up(ra.new_session())
            self.assertEqual(st, 200); self.assertTrue(body["ok"], body)
            ra.reset_state()


class AttachmentUiTests(unittest.TestCase):
    def test_composer_has_attach_paste_and_drop(self):
        html = (ROOT / "app_static" / "index.html").read_text(encoding="utf-8")
        js = (ROOT / "app_static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('id="attach"', html); self.assertIn('id="att-file"', html); self.assertIn("multiple", html)
        for want in ("/api/upload", "addEventListener('paste'", "addEventListener('drop'", "files: files.map"):
            self.assertIn(want, js)
        self.assertNotIn("base64,", js)
        self.assertNotIn("readAsDataURL", js)


if __name__ == "__main__":
    unittest.main()
