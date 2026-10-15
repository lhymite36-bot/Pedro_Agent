"""Chat attachments (1.4.9): pictures and files sent from the PC window or the phone app.

Files are saved on this PC only, in Documents\\PedroAgent\\uploads. Pictures and PDFs go to
Gemini as inline parts it can see/read; Word, Excel, CSV and text files go in as their text.
Nothing here sends anything anywhere else.
"""

from __future__ import annotations

import io
import json
import re
import time
import uuid
from pathlib import Path
from typing import Any

MAX_BYTES = 15 * 1024 * 1024
MAX_TEXT = 30_000
MAX_SIDE = 1600
IMAGE_EXT = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
             ".gif": "image/gif", ".bmp": "image/bmp", ".heic": "image/heic"}
DOC_EXT = {".pdf": "application/pdf", ".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "text",
           ".txt": "text", ".md": "text", ".json": "text", ".log": "text", ".py": "text", ".html": "text"}
ALLOWED = {**IMAGE_EXT, **DOC_EXT}
MARK = re.compile(r"\n?\[\[files:([0-9a-f,]{0,400})\]\]\s*$")
upload_dir_override: Path | None = None


def upload_dir() -> Path:
    d = upload_dir_override or (Path.home() / "Documents" / "PedroAgent" / "uploads")
    d.mkdir(parents=True, exist_ok=True)
    return d


def _index() -> dict[str, Any]:
    try:
        return json.loads((upload_dir() / "index.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def _save_index(data: dict[str, Any]) -> None:
    (upload_dir() / "index.json").write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")


def _clean_name(name: str) -> str:
    name = Path(str(name or "file").replace("\\", "/")).name
    name = re.sub(r"[^\w.\- ()]+", "_", name).strip(" .") or "file"
    return name[-120:]


def kind_of(name: str) -> str:
    return "image" if Path(name).suffix.lower() in IMAGE_EXT else "file"


def save_upload(name: str, data: bytes) -> dict[str, Any]:
    name = _clean_name(name)
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED:
        return {"ok": False, "error": f"{ext or 'That file type'} can't be attached. Use a picture, PDF, Word, Excel, CSV or text file."}
    if not data:
        return {"ok": False, "error": "The file is empty."}
    if len(data) > MAX_BYTES:
        return {"ok": False, "error": f"{name} is over {MAX_BYTES // (1024 * 1024)} MB."}
    if kind_of(name) == "image" and ext not in {".heic"}:
        try:
            from PIL import Image
            Image.open(io.BytesIO(data)).verify()
        except Exception:
            return {"ok": False, "error": f"{name} doesn't look like a real picture."}
    fid = uuid.uuid4().hex[:16]
    path = upload_dir() / f"{time.strftime('%Y%m%d-%H%M%S')}_{fid[:6]}_{name}"
    path.write_bytes(data)
    item = {"id": fid, "name": name, "kind": kind_of(name), "size": len(data), "path": str(path),
            "t": time.strftime("%Y-%m-%d %H:%M")}
    idx = _index(); idx[fid] = item; _save_index(idx)
    return {"ok": True, **item, "url": f"/api/upload?id={fid}"}


def get(fid: str) -> dict[str, Any] | None:
    if not re.fullmatch(r"[0-9a-f]{16}", str(fid or "")):
        return None
    item = _index().get(fid)
    if not item:
        return None
    p = Path(item["path"]).resolve()
    if upload_dir().resolve() not in p.parents or not p.is_file():
        return None
    return {**item, "path": str(p)}


def mime_for(name: str) -> str:
    ext = Path(name).suffix.lower()
    if ext in IMAGE_EXT:
        return IMAGE_EXT[ext]
    return {".pdf": "application/pdf", ".csv": "text/csv", ".txt": "text/plain", ".md": "text/plain",
            ".json": "application/json"}.get(ext, "application/octet-stream")


def _image_bytes(path: Path) -> tuple[bytes, str]:
    """Shrink big pictures so the model gets them fast; keeps small ones as they are."""
    data = path.read_bytes()
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        if max(im.size) <= MAX_SIDE and len(data) <= 3 * 1024 * 1024 and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
            return data, mime_for(path.name)
        im.thumbnail((MAX_SIDE, MAX_SIDE))
        out = io.BytesIO(); im.convert("RGB").save(out, "JPEG", quality=85)
        return out.getvalue(), "image/jpeg"
    except Exception:
        return data, mime_for(path.name)


def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if DOC_EXT.get(ext) == "docx":
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in r.cells))
            text = "\n".join(parts)
        elif DOC_EXT.get(ext) == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        parts.append(",".join("" if v is None else str(v) for v in row))
                    if sum(len(p) for p in parts) > MAX_TEXT:
                        break
            wb.close()
            text = "\n".join(parts)
        else:
            raw = path.read_bytes()
            text = raw.decode("utf-8", errors="replace")
    except Exception as exc:
        return f"(could not read {path.name}: {exc})"
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT] + f"\n... (cut; {len(text)} characters in total)"
    return text


def message_with_marker(text: str, ids: list[str]) -> str:
    ids = [i for i in ids if get(i)]
    return (text + ("\n" if text else "") + "[[files:" + ",".join(ids) + "]]") if ids else text


def split_marker(text: str) -> tuple[str, list[str]]:
    m = MARK.search(text or "")
    if not m:
        return text, []
    return text[:m.start()], [i for i in m.group(1).split(",") if i]


def build_turn(text: str, ids: list[str]) -> tuple[str, list[tuple[bytes, str]]]:
    """(text for the model, inline parts). Pictures/PDFs become parts; other files become text."""
    parts: list[tuple[bytes, str]] = []
    notes: list[str] = []
    for fid in ids:
        item = get(fid)
        if not item:
            continue
        p = Path(item["path"])
        if item["kind"] == "image":
            parts.append(_image_bytes(p)); notes.append(f"- picture {item['name']} (attached below; saved at {p})")
        elif p.suffix.lower() == ".pdf":
            parts.append((p.read_bytes(), "application/pdf")); notes.append(f"- PDF {item['name']} (attached below; saved at {p})")
        else:
            notes.append(f"- file {item['name']} (saved at {p}). Its content:\n<<<{item['name']}\n{read_text(p)}\n>>>")
    if not notes:
        return text, []
    head = (text or "Look at the attached files.") + "\n\n[Attached by the user in this message]\n" + "\n".join(notes)
    return head, parts


def info(ids: list[str]) -> list[dict[str, Any]]:
    out = []
    for fid in ids:
        it = get(fid)
        if it:
            out.append({"id": fid, "name": it["name"], "kind": it["kind"], "size": it["size"], "url": f"/api/upload?id={fid}"})
    return out
